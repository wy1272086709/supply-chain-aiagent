"""PPTX → Markdown — 自研 ZIP/XML 路径（逐字移植 pptx.parser.ts），失败降级 python-pptx。"""

from __future__ import annotations

import io
import re
import zipfile

from app.modules.documents.parsers.markdown_util import clean_markdown, to_markdown_table

_SLIDE_PATH_RE = re.compile(r"^ppt/slides/slide\d+\.xml$", re.IGNORECASE)
_SLIDE_NUM_RE = re.compile(r"slide(\d+)\.xml$", re.IGNORECASE)


def parse_pptx(data: bytes) -> str:
    try:
        return _parse_pptx_with_zip(data)
    except Exception:
        return _parse_pptx_with_python_pptx(data)


def _slide_number(path: str) -> int:
    m = _SLIDE_NUM_RE.search(path)
    return int(m.group(1)) if m else 0


def _decode_xml(text: str) -> str:
    return (
        text.replace("&lt;", "<")
        .replace("&gt;", ">")
        .replace("&quot;", '"')
        .replace("&apos;", "'")
        .replace("&amp;", "&")
    )


_RUN_RE = re.compile(r"<a:br\s*/>|<a:t(?:\s[^>]*)?>([\s\S]*?)</a:t>")


def _extract_runs_text(xml: str) -> str:
    result = ""
    for m in _RUN_RE.finditer(xml):
        if m.group(0).startswith("<a:br"):
            result += "\n"
        else:
            result += _decode_xml(m.group(1))
    return result


_P_BLOCK_RE = re.compile(r"<a:p[\s>][\s\S]*?</a:p>")


def _extract_paragraph_texts(xml: str) -> list[str]:
    paragraphs: list[str] = []
    for p in _P_BLOCK_RE.finditer(xml):
        text = _extract_runs_text(p.group(0))
        if text.strip():
            paragraphs.append(text)
    return paragraphs


_PH_RE = re.compile(r'<p:ph[^>]*\btype="([^"]+)"', re.IGNORECASE)


def _extract_placeholder_texts(xml: str, type_pattern: re.Pattern) -> list[str]:
    texts: list[str] = []
    shapes = xml.split("<p:sp ")[1:] + xml.split("<p:sp>")[1:]
    for shape in shapes:
        ph = _PH_RE.search(shape)
        if not ph or not type_pattern.search(ph.group(1)):
            continue
        joined = "\n".join(_extract_paragraph_texts(shape)).strip()
        if joined:
            texts.append(joined)
    return texts


_TBL_RE = re.compile(r"<a:tbl[\s\S]*?</a:tbl>")
_TR_RE = re.compile(r"<a:tr[\s\S]*?</a:tr>")
_TC_RE = re.compile(r"<a:tc[\s\S]*?</a:tc>")


def _extract_tables(xml: str) -> list[list[list[str]]]:
    tables: list[list[list[str]]] = []
    for block in _TBL_RE.findall(xml):
        rows: list[list[str]] = []
        for tr in _TR_RE.findall(block):
            cells = []
            for tc in _TC_RE.findall(tr):
                cell_paras = _extract_paragraph_texts(tc)
                cells.append(" ".join(cell_paras).strip())
            if cells:
                rows.append(cells)
        if rows:
            tables.append(rows)
    return tables


def _parse_pptx_with_zip(data: bytes) -> str:
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        slide_paths = sorted(
            [p for p in zf.namelist() if _SLIDE_PATH_RE.match(p)], key=_slide_number
        )
        if not slide_paths:
            raise ValueError("未找到 PPTX slide")
        parts: list[str] = []
        for i, path in enumerate(slide_paths):
            xml = zf.read(path).decode("utf-8", errors="replace")
            parts.append(f"## 幻灯片 {i + 1}\n")

            for table in _extract_tables(xml):
                parts.append(to_markdown_table(table))

            body_xml = re.sub(r"<a:tbl[\s\S]*?</a:tbl>", "", xml)

            title_texts = _extract_placeholder_texts(body_xml, re.compile(r"ctrTitle|title", re.IGNORECASE))
            body_paragraphs = _extract_paragraph_texts(body_xml)

            used = {t.strip() for t in title_texts if t.strip()}
            for t in title_texts:
                text = t.strip()
                if text:
                    parts.append(f"### {text}\n")

            body_lines = [
                t.strip() for t in body_paragraphs if t.strip() and t.strip() not in used
            ]
            if body_lines:
                parts.append("\n".join(body_lines) + "\n\n")

    result = clean_markdown("\n".join(parts))
    if not result:
        raise ValueError("PPTX 提取结果为空")
    return result


def _parse_pptx_with_python_pptx(data: bytes) -> str:
    """降级路径：python-pptx 结构化提取（对应 officeparser 降级）。"""
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts: list[str] = []
    for i, slide in enumerate(prs.slides):
        parts.append(f"## 幻灯片 {i + 1}\n")
        for shape in slide.shapes:
            if shape.has_table:
                rows = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                if rows:
                    parts.append(to_markdown_table(rows))
        for shape in slide.shapes:
            if getattr(shape, "is_placeholder", False):
                ph_type = str(shape.placeholder_format.type)
                if "TITLE" in ph_type.upper() and shape.has_text_frame:
                    text = shape.text_frame.text.strip()
                    if text:
                        parts.append(f"### {text}\n")
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = shape.text_frame.text.strip()
                if text:
                    parts.append(text + "\n\n")
    return clean_markdown("\n".join(parts))
