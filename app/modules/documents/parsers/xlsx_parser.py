"""XLSX → Markdown — openpyxl，对应 xlsx.parser.ts（exceljs）。

每个 Sheet → `## SheetName` + Markdown 表格（首行表头）；
富文本单元格 openpyxl 默认返回拼接后的纯文本；公式取缓存计算值。
"""

from __future__ import annotations

import io
from datetime import date, datetime, timedelta

from openpyxl import load_workbook

from app.modules.documents.parsers.markdown_util import clean_markdown, to_markdown_table


def _cell_to_string(value) -> str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "True" if value else "False"
    if isinstance(value, (int, float)):
        return str(value)
    if isinstance(value, datetime):
        return value.isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    # 富文本 / 超链接等对象
    text = getattr(value, "text", None)
    if isinstance(text, str):
        return text
    return str(value)


def parse_xlsx(data: bytes, data_only: bool = True) -> str:
    workbook = load_workbook(io.BytesIO(data), data_only=data_only, rich_text=False)
    parts: list[str] = []
    try:
        for sheet in workbook.worksheets:
            parts.append(f"## {sheet.title}\n")
            rows: list[list[str]] = []
            max_cols = 0
            for row in sheet.iter_rows():
                cells = [_cell_to_string(c.value) for c in row]
                if not any(cell.strip() for cell in cells):
                    continue  # includeEmpty: false
                max_cols = max(max_cols, len(cells))
                rows.append(cells)
            if max_cols == 0 or not rows:
                parts.append("\n")
                continue
            normalized = [r + [""] * (max_cols - len(r)) for r in rows]
            parts.append(to_markdown_table(normalized))
    finally:
        workbook.close()
    return clean_markdown("\n".join(parts))
